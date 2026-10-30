"""
Export Agent
============
Converts validated data into the requested file formats:
CSV, JSON, XLSX, XML, SQL (INSERT statements), Parquet, PDF, HTML, Markdown, and DOCX.
"""

from __future__ import annotations

import json
from html import escape
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import Element, SubElement, tostring

import pandas as pd

from app.agents.state import AgentState
from app.config import settings
from app.utils.logging_config import get_logger
from app.utils.security import check_export_file_size, safe_filename, sanitise_sql_value, validate_sql_safety

logger = get_logger(__name__)


def export_node(state: AgentState) -> dict[str, Any]:
    """LangGraph node — export validated data to the requested formats."""
    logger.info("[Export] Starting export for job %s", state.get("job_id"))

    data = state.get("generated_data")
    formats = state.get("requested_formats", ["csv"])
    job_id = state.get("job_id", "unknown")

    if not data:
        return {
            "status": "failed",
            "error": "Export agent received no data.",
            "current_agent": "export",
            "agent_messages": ["Export FAILED — no data to export."],
        }

    export_dir = Path(settings.EXPORT_DIR) / job_id
    export_dir.mkdir(parents=True, exist_ok=True)

    export_files: list[dict] = []
    messages: list[str] = []

    for fmt in formats:
        fmt_lower = fmt.lower().strip()
        try:
            for entity_name, records in data.items():
                if not records:
                    continue

                df = pd.DataFrame(records)
                safe_entity = safe_filename(entity_name)
                file_info = _export_entity(df, safe_entity, fmt_lower, export_dir, entity_name)

                if file_info:
                    # File-size guard
                    fpath = Path(file_info["path"])
                    if not check_export_file_size(fpath):
                        messages.append(f"WARNING: {file_info['filename']} exceeds size limit.")
                    export_files.append(file_info)
                    messages.append(f"Exported {entity_name} → {file_info['filename']}")

        except Exception as exc:
            logger.error("[Export] Failed to export format '%s': %s", fmt, exc)
            messages.append(f"Export error for format {fmt}: {exc}")

    logger.info("[Export] Completed — %d file(s) produced.", len(export_files))

    return {
        "export_files": export_files,
        "current_agent": "export",
        "status": "completed",
        "agent_messages": [f"Export completed — {len(export_files)} file(s)."] + messages,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Format-specific exporters
# ═══════════════════════════════════════════════════════════════════════════════

def _export_entity(df: pd.DataFrame, safe_name: str, fmt: str, export_dir: Path, entity_name: str) -> dict | None:
    """Export a single entity DataFrame to the given format. Returns file info dict."""
    version = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    versioned_name = f"{safe_name}_{version}"
    if fmt == "csv":
        file_info = _to_csv(df, versioned_name, export_dir, entity_name)
    elif fmt == "json":
        file_info = _to_json(df, versioned_name, export_dir, entity_name)
    elif fmt == "xlsx":
        file_info = _to_xlsx(df, versioned_name, export_dir, entity_name)
    elif fmt == "xml":
        file_info = _to_xml(df, versioned_name, export_dir, entity_name)
    elif fmt == "sql":
        file_info = _to_sql(df, versioned_name, export_dir, entity_name)
    elif fmt == "parquet":
        file_info = _to_parquet(df, versioned_name, export_dir, entity_name)
    elif fmt == "pdf":
        file_info = _to_pdf(df, versioned_name, export_dir, entity_name)
    elif fmt == "html":
        file_info = _to_html(df, versioned_name, export_dir, entity_name)
    elif fmt == "markdown":
        file_info = _to_markdown(df, versioned_name, export_dir, entity_name)
    elif fmt == "docx":
        file_info = _to_docx(df, versioned_name, export_dir, entity_name)
    else:
        logger.warning("[Export] Unsupported format: %s", fmt)
        return None

    if file_info:
        file_info["entity_name"] = entity_name
    return file_info


def _file_info(filename: str, fmt: str, path: Path, record_count: int) -> dict:
    return {
        "filename": filename,
        "format": fmt,
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "record_count": record_count,
        "created_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
    }


def _to_csv(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.csv"
    path = out / filename
    df.to_csv(path, index=False, encoding="utf-8")
    return _file_info(filename, "csv", path, len(df))


def _to_json(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.json"
    path = out / filename
    records = df.to_dict(orient="records")
    path.write_text(json.dumps(records, indent=2, default=str), encoding="utf-8")
    return _file_info(filename, "json", path, len(df))


def _to_xlsx(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.xlsx"
    path = out / filename
    df.to_excel(path, index=False, sheet_name=entity[:31], engine="openpyxl")
    return _file_info(filename, "xlsx", path, len(df))


def _to_xml(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.xml"
    path = out / filename

    root = Element("dataset")
    root.set("entity", entity)
    for _, row in df.iterrows():
        record_el = SubElement(root, "record")
        for col in df.columns:
            field_el = SubElement(record_el, safe_filename(str(col)))
            val = row[col]
            field_el.text = "" if pd.isna(val) else str(val)

    xml_bytes = tostring(root, encoding="unicode", xml_declaration=False)
    path.write_text(f'<?xml version="1.0" encoding="UTF-8"?>\n{xml_bytes}', encoding="utf-8")
    return _file_info(filename, "xml", path, len(df))


def _to_sql(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.sql"
    path = out / filename
    table_name = safe_filename(entity)
    lines: list[str] = [f"-- Generated INSERT statements for table: {table_name}", ""]

    for _, row in df.iterrows():
        cols = ", ".join(f'"{c}"' for c in df.columns)
        vals: list[str] = []
        for val in row:
            if pd.isna(val) or val is None:
                vals.append("NULL")
            elif isinstance(val, (int, float)):
                vals.append(str(val))
            elif isinstance(val, bool):
                vals.append("TRUE" if val else "FALSE")
            else:
                vals.append(f"'{sanitise_sql_value(str(val))}'")
        values_str = ", ".join(vals)
        line = f'INSERT INTO "{table_name}" ({cols}) VALUES ({values_str});'
        lines.append(line)

    sql_content = "\n".join(lines)

    # Safety check
    if not validate_sql_safety(sql_content):
        logger.error("[Export] SQL export blocked by safety check for entity %s.", entity)
        sql_content = f"-- SQL export blocked by safety check for entity: {entity}\n"

    path.write_text(sql_content, encoding="utf-8")
    return _file_info(filename, "sql", path, len(df))


def _to_parquet(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.parquet"
    path = out / filename
    df.to_parquet(path, index=False, engine="pyarrow")
    return _file_info(filename, "parquet", path, len(df))


def _to_pdf(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    filename = f"{name}.pdf"
    path = out / filename
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        str(path),
        pagesize=page_size,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=entity,
    )

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        "TableHeader",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )
    if _is_financial_statement(df):
        story = _build_financial_pdf_story(df, entity, styles, header_style, cell_style)
        document.build(story)
        return _file_info(filename, "pdf", path, len(df))

    headers = [Paragraph(escape(str(column)), header_style) for column in df.columns]
    rows = [
        [
            Paragraph("" if pd.isna(value) else escape(str(value)), cell_style)
            for value in row
        ]
        for row in df.itertuples(index=False, name=None)
    ]
    if not headers:
        headers = [Paragraph("(no fields)", header_style)]
        rows = [[""] for _ in range(max(len(df), 1))]
    column_count = max(len(headers), 1)
    available_width = page_size[0] - document.leftMargin - document.rightMargin
    table = LongTable(
        [headers, *rows],
        colWidths=[available_width / column_count] * column_count,
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312e81")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    document.build([
        Paragraph(escape(entity), title_style),
        Paragraph(f"{len(df)} records", styles["Normal"]),
        Spacer(1, 8),
        table,
    ])
    return _file_info(filename, "pdf", path, len(df))


def _is_financial_statement(df: pd.DataFrame) -> bool:
    columns = {_normalize_report_label(str(column)) for column in df.columns}
    has_line_item = bool(columns & {"lineitem", "account", "accountname"})
    has_amount = bool(columns & {"amount", "balance", "value", "currentperiodamount"})
    return has_line_item and has_amount


def _normalize_report_label(value: str) -> str:
    return "".join(character.lower() for character in value if character.isalnum())


def _build_financial_pdf_story(df, entity, styles, header_style, cell_style):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    normalized_columns = {
        _normalize_report_label(str(column)): column for column in df.columns
    }

    def find_column(*names):
        return next(
            (normalized_columns[name] for name in names if name in normalized_columns),
            None,
        )

    line_column = find_column("lineitem", "accountname", "account")
    amount_column = find_column("amount", "balance", "value", "currentperiodamount")
    section_column = find_column("statementsection", "section", "category")
    period_column = find_column("period", "fiscalperiod", "fiscalyear", "date")
    currency_column = find_column("currency")

    periods: list[str] = []
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    section_order: list[str] = []
    for _, record in df.iterrows():
        section = str(record.get(section_column) or "Statement details") if section_column else "Statement details"
        line_item = str(record.get(line_column) or "Unlabeled item")
        period = str(record.get(period_column) or "Amount") if period_column else "Amount"
        if period not in periods:
            periods.append(period)
        if section not in section_order:
            section_order.append(section)
        key = (section, line_item)
        if key not in grouped:
            grouped[key] = {"amounts": {}, "currency": record.get(currency_column) if currency_column else None}
        grouped[key]["amounts"][period] = record.get(amount_column)

    title_style = styles["Title"]
    subtitle_style = ParagraphStyle(
        "FinancialSubtitle",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.HexColor("#475569"),
    )
    story = [
        Paragraph(escape(entity), title_style),
        Paragraph("Synthetic scenario data — not an actual company filing.", subtitle_style),
        Spacer(1, 8),
    ]

    table_rows = [[
        Paragraph("Line item", header_style),
        *[Paragraph(escape(period), header_style) for period in periods],
    ]]
    section_row_indices: list[int] = []
    total_row_indices: list[int] = []
    for section in section_order:
        section_row_indices.append(len(table_rows))
        table_rows.append([Paragraph(escape(section), header_style)] + [""] * len(periods))
        for current_section, line_item in grouped:
            if current_section != section:
                continue
            item = grouped[(current_section, line_item)]
            currency = item["currency"]
            values = [
                _format_report_amount(item["amounts"].get(period), currency)
                for period in periods
            ]
            row_index = len(table_rows)
            table_rows.append([
                Paragraph(escape(line_item), cell_style),
                *[
                    Paragraph(escape(value), cell_style)
                    for value in values
                ],
            ])
            if any(marker in _normalize_report_label(line_item) for marker in (
                "total", "grossprofit", "operatingincome", "netincome", "netloss",
            )):
                total_row_indices.append(row_index)

    page_width = 297 * mm
    left_right_margins = 30 * mm
    available_width = page_width - left_right_margins
    label_width = available_width * 0.52
    amount_width = (available_width - label_width) / max(len(periods), 1)
    table = Table(
        table_rows,
        colWidths=[label_width] + [amount_width] * len(periods),
        repeatRows=1,
        hAlign="LEFT",
    )
    table_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312e81")),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#312e81")),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for row_index in section_row_indices:
        table_style.extend([
            ("SPAN", (0, row_index), (-1, row_index)),
            ("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#e2e8f0")),
            ("TEXTCOLOR", (0, row_index), (-1, row_index), colors.HexColor("#0f172a")),
            ("LINEABOVE", (0, row_index), (-1, row_index), 0.6, colors.HexColor("#94a3b8")),
        ])
    for row_index in total_row_indices:
        table_style.extend([
            ("FONTNAME", (0, row_index), (-1, row_index), "Helvetica-Bold"),
            ("LINEABOVE", (0, row_index), (-1, row_index), 0.6, colors.HexColor("#64748b")),
        ])
    table.setStyle(TableStyle(table_style))
    story.append(table)
    return story


def _format_report_amount(value: Any, currency: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    try:
        amount = float(str(value).replace(",", "").replace("$", ""))
        rendered = f"({abs(amount):,.2f})" if amount < 0 else f"{amount:,.2f}"
    except (ValueError, TypeError):
        rendered = str(value)
    return f"{currency} {rendered}" if currency else rendered


def _to_html(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.html"
    path = out / filename
    table = df.to_html(index=False, escape=True, border=0, classes="data-table")
    content = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{escape(entity)}</title>
  <style>
    body {{ color: #0f172a; font: 14px/1.5 Arial, sans-serif; margin: 2rem; }}
    h1 {{ color: #312e81; }}
    .table-wrap {{ overflow-x: auto; }}
    .data-table {{ border-collapse: collapse; width: 100%; }}
    .data-table th, .data-table td {{ border: 1px solid #cbd5e1; padding: .5rem; text-align: left; }}
    .data-table th {{ background: #312e81; color: white; }}
    .data-table tbody tr:nth-child(even) {{ background: #f1f5f9; }}
  </style>
</head>
<body>
  <h1>{escape(entity)}</h1>
  <p>{len(df)} records</p>
  <div class="table-wrap">{table}</div>
</body>
</html>
"""
    path.write_text(content, encoding="utf-8")
    return _file_info(filename, "html", path, len(df))


def _to_markdown(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    filename = f"{name}.md"
    path = out / filename

    def format_cell(value: Any) -> str:
        if pd.isna(value):
            return ""
        return str(value).replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")

    headers = [format_cell(column) for column in df.columns]
    lines = [
        f"# {entity}",
        "",
        f"{len(df)} records",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend(
        "| " + " | ".join(format_cell(value) for value in row) + " |"
        for row in df.itertuples(index=False, name=None)
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return _file_info(filename, "markdown", path, len(df))


def _to_docx(df: pd.DataFrame, name: str, out: Path, entity: str) -> dict:
    from docx import Document

    filename = f"{name}.docx"
    path = out / filename
    document = Document()
    document.add_heading(entity, level=1)
    document.add_paragraph(f"{len(df)} records")

    columns = [str(column) for column in df.columns] or ["(no fields)"]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Light Shading Accent 1"
    for cell, column in zip(table.rows[0].cells, columns):
        cell.text = column

    if df.columns.empty:
        values = [[""] for _ in range(max(len(df), 1))]
    else:
        values = [
            ["" if pd.isna(value) else str(value) for value in row]
            for row in df.itertuples(index=False, name=None)
        ]
    for row in values:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value

    document.save(path)
    return _file_info(filename, "docx", path, len(df))
